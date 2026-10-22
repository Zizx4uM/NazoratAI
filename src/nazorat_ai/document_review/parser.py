"""
Multi-format document parser for NazoratAI.
Supports .docx, .pdf, .xlsx, .xls, and .txt files.
Extracts clean text, tabular data, and key regulatory monitoring fields.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Union
import re
import pandas as pd

try:
    import docx
except ImportError:
    docx = None

try:
    import pypdf
except ImportError:
    pypdf = None


@dataclass
class ParsedDocument:
    """Standardized parsed document representation."""
    file_path: str
    file_name: str
    file_type: str
    raw_text: str
    cleaned_text: str
    tables: List[List[List[str]]] = field(default_factory=list)
    char_count: int = 0
    extracted_inn: Optional[str] = None
    extracted_company: Optional[str] = None
    extracted_date: Optional[str] = None


class DocumentParser:
    """Multi-format file parser for credit monitoring acts and reports."""

    def parse(self, file_path: Union[str, Path]) -> ParsedDocument:
        """Parse file from path and return structured ParsedDocument."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Document file not found: {path}")

        suffix = path.suffix.lower()
        if suffix == ".docx":
            return self._parse_docx(path)
        elif suffix == ".pdf":
            return self._parse_pdf(path)
        elif suffix in (".xlsx", ".xls"):
            return self._parse_excel(path)
        elif suffix == ".txt":
            return self._parse_text(path)
        else:
            raise ValueError(f"Unsupported document format: {suffix}. Supported: .docx, .pdf, .xlsx, .txt")

    def _parse_docx(self, path: Path) -> ParsedDocument:
        if docx is None:
            raise RuntimeError("python-docx is not installed.")

        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        raw_text = "\n".join(paragraphs)

        tables_data: List[List[List[str]]] = []
        table_text_pieces: List[str] = []
        for table in doc.tables:
            t_rows: List[List[str]] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                t_rows.append(cells)
                table_text_pieces.append(" | ".join(cells))
            tables_data.append(t_rows)

        full_text = raw_text + "\n" + "\n".join(table_text_pieces)
        return self._build_parsed_document(path, full_text, tables_data, "docx")

    def _parse_pdf(self, path: Path) -> ParsedDocument:
        if pypdf is None:
            raise RuntimeError("pypdf is not installed.")

        reader = pypdf.PdfReader(str(path))
        text_pieces = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_pieces.append(extracted)

        full_text = "\n".join(text_pieces)
        return self._build_parsed_document(path, full_text, [], "pdf")

    def _parse_excel(self, path: Path) -> ParsedDocument:
        excel_file = pd.ExcelFile(str(path))
        text_pieces = []
        tables_data: List[List[List[str]]] = []

        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            text_pieces.append(f"--- Sheet: {sheet_name} ---")
            sheet_rows = [df.columns.astype(str).tolist()] + df.fillna("").astype(str).values.tolist()
            tables_data.append(sheet_rows)
            for row in sheet_rows:
                text_pieces.append(" | ".join(row))

        full_text = "\n".join(text_pieces)
        clean_ext = path.suffix.lower().lstrip(".")
        return self._build_parsed_document(path, full_text, tables_data, clean_ext)

    def _parse_text(self, path: Path) -> ParsedDocument:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            full_text = f.read()
        return self._build_parsed_document(path, full_text, [], "txt")

    def _build_parsed_document(
        self,
        path: Path,
        raw_text: str,
        tables: List[List[List[str]]],
        file_type: str,
    ) -> ParsedDocument:
        cleaned_text = re.sub(r"\s+", " ", raw_text).strip()

        # Regex extract INN (9 digits)
        inn_match = re.search(r"\b(?:INN|STIR|inn|stir)[:\s]*([0-9]{9})\b", raw_text, re.IGNORECASE)
        inn = inn_match.group(1) if inn_match else None
        if not inn:
            standalone_9d = re.search(r"\b([1-3][0-9]{8})\b", raw_text)
            if standalone_9d:
                inn = standalone_9d.group(1)

        # Regex extract Company Name
        company = None
        comp_match = re.search(
            r"(?:korxona|mijoz|qarz\s+oluvchi|nomi)[:\s]*[\"']?([A-Za-z0-9\s'ʻ`\-]+(?:MCHJ|AJ|XK|QK))[\"']?",
            raw_text,
            re.IGNORECASE,
        )
        if comp_match:
            company = comp_match.group(1).strip()

        # Regex extract Date
        date_match = re.search(r"\b(202[0-9][\-\/.][0-1][0-9][\-\/.][0-3][0-9])\b", raw_text)
        doc_date = date_match.group(1) if date_match else None

        return ParsedDocument(
            file_path=str(path),
            file_name=path.name,
            file_type=file_type,
            raw_text=raw_text,
            cleaned_text=cleaned_text,
            tables=tables,
            char_count=len(raw_text),
            extracted_inn=inn,
            extracted_company=company,
            extracted_date=doc_date,
        )
