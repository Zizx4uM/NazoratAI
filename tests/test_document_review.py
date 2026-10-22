"""
Tests for document parsing and hybrid document review (offline heuristics & regex NLP).
"""

import tempfile
from pathlib import Path
import pytest
import docx
import pandas as pd

from nazorat_ai.config import RiskTier
from nazorat_ai.document_review.parser import DocumentParser
from nazorat_ai.document_review.analyzer import HybridDocumentReviewer


@pytest.fixture
def parser():
    return DocumentParser()


@pytest.fixture
def reviewer():
    return HybridDocumentReviewer()


def test_txt_parsing_and_extraction(parser):
    """Test text file parsing and regex entity extraction."""
    content = (
        "KREDIT MONITORINGI DALOLATNOMASI\n"
        "Korxona: 'Samarqand Shisha Zavodi' AJ\n"
        "STIR (INN): 301234567\n"
        "Sana: 2026-09-15\n"
        "Obyekt ko'zdan kechirildi. Ishlab chiqarish reja asosida bormoqda."
    )
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", encoding="utf-8", delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)

    doc = parser.parse(tmp_path)
    assert doc.file_type == "txt"
    assert doc.extracted_inn == "301234567"
    assert "Samarqand Shisha Zavodi" in (doc.extracted_company or "")
    assert doc.extracted_date == "2026-09-15"


def test_docx_parsing(parser):
    """Test .docx document parsing."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc = docx.Document()
        doc.add_heading("Monitoring Dalolatnomasi", level=1)
        doc.add_paragraph("Mijoz: 'O'zbekiston Tekstil Klasteri' MCHJ")
        doc.add_paragraph("INN: 309876543")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Ko'rsatkich"
        table.cell(0, 1).text = "Holat"
        table.cell(1, 0).text = "Liniya"
        table.cell(1, 1).text = "Ishlamoqda"
        doc.save(tmp.name)
        tmp_path = Path(tmp.name)

    parsed = parser.parse(tmp_path)
    assert parsed.file_type == "docx"
    assert parsed.extracted_inn == "309876543"
    assert len(parsed.tables) == 1


def test_excel_parsing(parser):
    """Test Excel (.xlsx) parsing."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        df = pd.DataFrame({
            "INN": ["304555666"],
            "Korxona": ["Buxoro Limonariya XK"],
            "Qoldiq": [1500000000],
        })
        df.to_excel(tmp.name, index=False)
        tmp_path = Path(tmp.name)

    parsed = parser.parse(tmp_path)
    assert parsed.file_type == "xlsx"
    assert "304555666" in parsed.raw_text


def test_offline_heuristic_review_clean_doc(reviewer):
    """Clean monitoring act with no red flags should yield GREEN risk tier."""
    content = (
        "MONITORING DALOLATNOMASI\n"
        "Korxona: 'Ideal Qurilish' MCHJ\n"
        "INN: 301111222\n"
        "Korxona to'liq quvvat bilan faoliyat yuritmoqda. "
        "Uskunalar soz holatda, mahsulot ishlab chiqarish grafigi bajarilmoqda. "
        "Garov mulki butun va amalda mavjud."
    )
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", encoding="utf-8", delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)

    result = reviewer.review_file(tmp_path)
    assert result.overall_risk_level == RiskTier.GREEN.value
    assert result.flags_count == 0
    assert result.audit_mode == "LOCAL_OFFLINE_HEURISTIC"


def test_offline_heuristic_review_critical_red_flags(reviewer):
    """Text mentioning factory shutdown and collateral missing triggers RED."""
    content = (
        "FAVQULODDA MONITORING DALOLATNOMASI\n"
        "Korxona: 'Inqiroz Tekstil' MCHJ\n"
        "INN: 302222333\n"
        "Joyiga chiqib o'rganganda, sexlar ishlamayapti, ishlab chiqarish to'xtatilgan. "
        "Elektr uzilgan sababli liniyalar to'xtab qolgan. "
        "Bundan tashqari, garov mulki sotilganligi va joyida topilmadi. "
        "Korxonada ishchilar qisqartirilgan va 2-kartoteka mavjud."
    )
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", encoding="utf-8", delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)

    result = reviewer.review_file(tmp_path)
    assert result.overall_risk_level == RiskTier.RED.value
    assert result.critical_flags_count >= 1
    flag_ids = [f.flag_id for f in result.flags]
    assert "DOC-FLAG-01" in flag_ids  # Ishlab chiqarish to'xtagan
    assert "DOC-FLAG-03" in flag_ids  # Garov sotilgan
    assert "DOC-FLAG-04" in flag_ids  # Elektr uzilgan
    assert "DOC-FLAG-05" in flag_ids  # Ishchilar qisqartirilgan
    assert "DOC-FLAG-07" in flag_ids  # 2-kartoteka


def test_offline_heuristic_empty_document(reviewer):
    """Empty or unreadable document should return UNKNOWN rather than falsely claiming GREEN."""
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", encoding="utf-8", delete=False) as tmp:
        tmp.write("   \n\n  \t  ")
        tmp_path = Path(tmp.name)

    result = reviewer.review_file(tmp_path)
    assert result.overall_risk_level == "UNKNOWN"
    assert "bo'sh" in result.executive_summary.lower()
